"""Artifact assertions beyond geometry (M5+): drawings, DXF, BOM files, docs, report."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

EXTERNAL = re.compile(r"\b(?:src|href)\s*=\s*['\"]?(?:https?:)?//", re.IGNORECASE)
MAX_REPORT = 20 * 1024 * 1024


def validate_full(root: Path, product, problems: list[str]) -> None:  # type: ignore[no-untyped-def]
    """Drawings for every custom part, DXF for sheet metal, BOM CSV/XLSX, docs HTML/PDF, report."""
    import ezdxf
    import openpyxl
    from lxml import etree
    from selectolax.lexbor import LexborHTMLParser

    build = root / "build"
    # drawings
    for part in product.parts_with_geometry():
        svg = build / "drawings" / f"{part.id}.svg"
        if not svg.is_file():
            problems.append(f"{part.uri}: missing drawing {svg.name}")
            continue
        try:
            tree = etree.parse(str(svg))
        except etree.XMLSyntaxError as exc:
            problems.append(f"{part.uri}: drawing SVG does not parse: {exc}")
            continue
        ids = {el.get("id") for el in tree.getroot().iter() if el.get("id")}
        for g in ("view-front", "view-top", "view-right", "view-iso", "titleblock"):
            if g not in ids:
                problems.append(f"{part.uri}: drawing lacks <g id='{g}'>")
        if part.process == "sheet-metal":
            dxf = build / "drawings" / f"{part.id}-flat.dxf"
            if not dxf.is_file():
                problems.append(f"{part.uri}: sheet-metal part has no flat-pattern DXF")
            else:
                doc = ezdxf.readfile(str(dxf))
                layers = {ln.dxf.name for ln in doc.layers}
                if "BEND" not in layers:
                    problems.append(f"{part.uri}: DXF has no BEND layer")
            if "view-flat" not in ids:
                problems.append(f"{part.uri}: sheet-metal drawing has no flat-pattern view")
    # bom
    for name in ("rollups.json", "validation.json", "bom.csv", "bom.xlsx"):
        if not (build / "bom" / name).is_file():
            problems.append(f"bom: missing {name}")
    rollups = (
        json.loads((build / "bom" / "rollups.json").read_text())
        if (build / "bom" / "rollups.json").is_file()
        else {}
    )
    if (build / "bom" / "bom.xlsx").is_file():
        wb = openpyxl.load_workbook(str(build / "bom" / "bom.xlsx"), data_only=True)
        if "BOM" not in wb.sheetnames:
            problems.append("bom.xlsx has no BOM sheet")
        else:
            numbers = [
                float(c.value)
                for row in wb["BOM"].iter_rows()
                for c in row
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
            ]
            if rollups and not any(abs(n - float(rollups["total_cost"])) < 0.005 for n in numbers):
                problems.append("bom.xlsx totals do not include rollups total_cost")
    # build plan (P8): one step per build unit, its sheet in the workbook, one assembly
    # sheet per unit and none for a phantom
    from napkin2factory.core.build_units import BuildUnits

    units = BuildUnits(product).order()
    plan_path = build / "bom" / "build-plan.json"
    if product.assemblies and not plan_path.is_file():
        problems.append("bom: missing build-plan.json")
    elif plan_path.is_file():
        plan = json.loads(plan_path.read_text())
        if plan.get("order") != units:
            problems.append(f"build-plan.json steps {plan.get('order')} != build units {units}")
        if (build / "bom" / "bom.xlsx").is_file():
            wb = openpyxl.load_workbook(str(build / "bom" / "bom.xlsx"), data_only=True)
            if "Build plan" not in wb.sheetnames:
                problems.append("bom.xlsx has no Build plan sheet")
    for aid, asm in product.assemblies.items():
        sheet = build / "drawings" / "assemblies" / f"{aid}.svg"
        if aid in units and not sheet.is_file():
            problems.append(f"{asm.uri}: build unit without an assembly sheet")
        elif aid not in units and sheet.exists():
            problems.append(f"{asm.uri}: a phantom has an assembly sheet")
        elif sheet.is_file():
            try:
                root_el = etree.parse(str(sheet)).getroot()
            except etree.XMLSyntaxError as exc:
                problems.append(f"{asm.uri}: assembly sheet does not parse: {exc}")
                continue
            ids = {el.get("id") for el in root_el.iter() if el.get("id")}
            if not {"pictures", "tables", "titleblock"} <= ids:
                problems.append(f"{asm.uri}: assembly sheet lacks pictures/tables/titleblock")
    if (build / "bom" / "bom.csv").is_file():
        with (build / "bom" / "bom.csv").open(newline="", encoding="utf-8") as fh:
            rows = list(csv.reader(fh))
        if not rows or "uri" not in rows[0]:
            problems.append("bom.csv has no header with 'uri'")
    # docs
    for doc in product.docs.values():
        d = build / "docs" / doc.meta.id
        for name in ("page.md", "page.html", "reads.json"):
            if not (d / name).is_file():
                problems.append(f"{doc.uri}: missing {name}")
        pdf = d / "page.pdf"
        if not pdf.is_file():
            problems.append(
                f"{doc.uri}: missing page.pdf"
                + (
                    f" ({(d / 'page.pdf.error.txt').read_text()[:120]})"
                    if (d / "page.pdf.error.txt").is_file()
                    else ""
                )
            )
        else:
            data = pdf.read_bytes()
            if data[:5] != b"%PDF-" or b"%%EOF" not in data[-64:] or len(data) < 500:
                problems.append(f"{doc.uri}: page.pdf is not a complete PDF")
        if (d / "page.html").is_file():
            tree = LexborHTMLParser((d / "page.html").read_text(encoding="utf-8"))
            if tree.css_first("h1") is None and tree.css_first("h2") is None:
                problems.append(f"{doc.uri}: page.html has no headings")
    safety_pages = [d for d in product.docs.values() if d.meta.safety]
    if not safety_pages:
        problems.append("docs: no page with safety: true (a Safety section is mandatory)")
    # checks
    checks = build / "checks" / "checks.json"
    if not checks.is_file():
        problems.append("checks: missing checks.json")
    else:
        data = json.loads(checks.read_text())
        fails = [r["id"] for r in data.get("results", []) if r["status"] == "fail"]
        if fails:
            problems.append(f"checks: failing rows: {', '.join(fails[:8])}")
    # report
    report = build / "report" / "index.html"
    if not report.is_file():
        problems.append("report: missing build/report/index.html")
    else:
        html = report.read_text(encoding="utf-8")
        size = report.stat().st_size
        if size > MAX_REPORT:
            problems.append(f"report: {size / 1e6:.1f} MB exceeds the 20 MB budget (D13)")
        tree = LexborHTMLParser(html)
        linked = tree.css('link:not([href^="data:"])')  # the embedded icon is a data URI
        if tree.css("script[src]") or linked or EXTERNAL.search(html):
            problems.append("report: has external script/link/asset references (D13)")
        modules = [s for s in tree.css("script") if s.attributes.get("type") == "module"]
        if len(modules) != 1 and any(p.has_geometry for p in product.parts.values()):
            problems.append(f"report: expected one inline module script, found {len(modules)}")
        for part in product.parts.values():
            if part.uri not in html:
                problems.append(f"report: {part.uri} not mentioned")
        for rid in product.requirements:
            if f'id="req-{rid}"' not in html:
                problems.append(f"report: requirement {rid} row missing")
        for aid in units:
            if f'id="step-{aid}"' not in html:
                problems.append(f"report: build plan step for asm://{aid} missing")
        if "not been reviewed by a licensed engineer" not in html:
            problems.append("report: disclaimer missing")
        for s in product.signoffs.values():
            if s.example and "EXAMPLE SIGN-OFF, NOT A REVIEW" not in html:
                problems.append(f"report: example sign-off {s.requirement} is not labelled")
    site = build / "site" / "index.html"
    if not site.is_file():
        problems.append("site: missing build/site/index.html")
