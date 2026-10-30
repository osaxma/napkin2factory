"""M0 toolchain spike: every third-party capability v0.1 depends on, proven once.

Each test is small and names the capability it proves so a failure after a dependency
bump points straight at the broken layer.
"""

from __future__ import annotations

from pathlib import Path

import pytest


@pytest.fixture(scope="module")
def box_with_hole():  # type: ignore[no-untyped-def]
    from build123d import Box, BuildPart, Hole, Locations

    with BuildPart() as bp:
        Box(40, 30, 20)
        with Locations((0, 0, 10)):
            Hole(5)
    part = bp.part
    part.label = "part://box-with-hole"
    return part


def test_build123d_box_with_hole(box_with_hole) -> None:  # type: ignore[no-untyped-def]
    import math

    expected = 40 * 30 * 20 - math.pi * 5**2 * 20
    assert abs(box_with_hole.volume - expected) < 1.0
    bb = box_with_hole.bounding_box().size
    assert (round(bb.X), round(bb.Y), round(bb.Z)) == (40, 30, 20)


def test_step_export_and_reimport(box_with_hole, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    from build123d import export_step, import_step

    out = tmp_path / "box.step"
    assert export_step(box_with_hole, str(out))
    back = import_step(str(out))
    assert abs(back.volume - box_with_hole.volume) < 1.0


def test_gltf_export_propagates_node_names(box_with_hole, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    """build123d's export_gltf carries ``label`` into node and mesh names, but writes no
    materials and does not share meshes between instances. The in-house exporter
    (pygltflib post-pass) therefore remains mandatory for D13; this test documents the
    baseline it builds on."""
    import pygltflib
    from build123d import export_gltf

    out = tmp_path / "box.glb"
    assert export_gltf(box_with_hole, str(out), binary=True)
    g = pygltflib.GLTF2().load(str(out))
    assert [n.name for n in g.nodes] == ["part://box-with-hole"]
    assert g.materials in (None, [])  # hence the post-pass


def test_svg_projection_has_hidden_lines(box_with_hole, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    from build123d import ExportSVG, LineType
    from lxml import etree

    visible, hidden = box_with_hole.project_to_viewport((100, -100, 80), viewport_up=(0, 0, 1))
    assert len(visible) > 0 and len(hidden) > 0
    exp = ExportSVG(scale=1)
    exp.add_layer("visible", line_weight=0.5)
    exp.add_layer("hidden", line_weight=0.25, line_type=LineType.ISO_DASH)
    exp.add_shape(visible, layer="visible")
    exp.add_shape(hidden, layer="hidden")
    out = tmp_path / "box.svg"
    exp.write(str(out))
    root = etree.parse(str(out)).getroot()
    ids = {el.get("id") for el in root.iter() if el.get("id")}
    assert {"visible", "hidden"} <= ids


def test_dxf_roundtrip(box_with_hole, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    import ezdxf
    from build123d import Axis, ExportDXF

    exp = ExportDXF()
    exp.add_layer("outline")
    exp.add_shape(box_with_hole.faces().sort_by(Axis.Z)[-1].outer_wire(), layer="outline")
    out = tmp_path / "box.dxf"
    exp.write(str(out))
    doc = ezdxf.readfile(str(out))
    assert len(list(doc.modelspace())) >= 4


def test_typst_compiles_pdf() -> None:
    import typst

    pdf = typst.compile(b"= Hello\nA paragraph with *emphasis*.")
    assert pdf[:5] == b"%PDF-"


def test_xlsx_roundtrip(tmp_path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["part", "qty"])
    wb.active.append(["part://x", 3])
    out = tmp_path / "bom.xlsx"
    wb.save(str(out))
    ws = openpyxl.load_workbook(str(out)).active
    assert ws["B2"].value == 3


def test_render_path_numpy_zbuffer(box_with_hole, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    """D12: the chosen wheel (novtk) has no VTK, so the z-buffer rasterizer is the path."""
    import numpy as np

    from napkin2factory.geometry.render import (
        Mesh,
        Scene,
        color_for,
        image_variance,
        rasterize,
        write_png,
    )

    verts, tris = box_with_hole.tessellate(0.2)
    mesh = Mesh(
        vertices=np.asarray([(v.X, v.Y, v.Z) for v in verts]),
        triangles=np.asarray(tris, dtype=np.int64),
        color=color_for("part://box-with-hole"),
    )
    img = rasterize(Scene([mesh]), "iso", size=(320, 240))
    assert img.shape == (240, 320, 3)
    assert image_variance(img) > 100.0, "render is blank"
    out = write_png(tmp_path / "iso.png", img)
    assert out.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"


def test_vtk_is_absent_as_recorded_in_decisions() -> None:
    """If VTK ever appears (wheel change), D2/D12 must be revisited in DECISIONS.md."""
    import importlib.util

    assert importlib.util.find_spec("vtk") is None
