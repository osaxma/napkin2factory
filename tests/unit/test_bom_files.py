"""BOM files (M5): bom.csv and bom.xlsx agree with rollups.json for examples/enclosure."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import pytest

from napkin2factory.bom.files import COLUMNS, SOURCING_NOTE
from tests.conftest import EXAMPLES, copy_product


@pytest.fixture(scope="module")
def built(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """A temp copy of examples/enclosure built to the ``bom`` target."""
    from napkin2factory.core.graph.runner import run_build
    from napkin2factory.core.loader import load_product

    root = copy_product(
        EXAMPLES / "enclosure", tmp_path_factory.mktemp("bom") / "enclosure", git=False
    )
    run_build(load_product(root), target="bom", quiet=True, render=False)
    return root / "build" / "bom"


def _rollups(bom_dir: Path) -> dict[str, Any]:
    return json.loads((bom_dir / "rollups.json").read_text(encoding="utf-8"))


def test_files_exist(built: Path) -> None:
    for name in ("validation.json", "bom.csv", "bom.xlsx", "rollups.json"):
        assert (built / name).is_file(), name


def test_csv(built: Path) -> None:
    text = (built / "bom.csv").read_text(encoding="utf-8")
    table, _, totals = text.partition("\n\n")
    rows = list(csv.DictReader(table.splitlines()))
    assert list(rows[0].keys()) == COLUMNS
    assert [r["id"] for r in rows if r["kind"] != "assembly"] == [
        "base",
        "lid",
        "bracket",
        "m4x10-screw",
    ]
    bracket = next(r for r in rows if r["id"] == "bracket")
    assert float(bracket["unit_cost_usd"]) == 4.00
    assert float(bracket["ext_cost_usd"]) == pytest.approx(2 * 4.00)
    assert int(bracket["total_qty"]) == 2
    base = next(r for r in rows if r["id"] == "base")
    assert base["uri"] == "part://base" and base["manufacturer"]

    tot = dict(r for r in csv.reader(totals.splitlines()) if r)
    rollups = _rollups(built)
    assert list(tot) == [
        "total_cost_usd",
        "total_mass_kg",
        "part_count",
        "custom_count",
        "custom_ratio",
        "max_lead_time_days",
        "cost_complete",
        "mass_complete",
    ]
    assert float(tot["total_cost_usd"]) == rollups["total_cost"]
    assert float(tot["total_mass_kg"]) == rollups["total_mass"]
    assert int(tot["part_count"]) == rollups["part_count"]
    assert tot["cost_complete"] == str(rollups["cost_complete"]).lower()


def test_xlsx(built: Path) -> None:
    import openpyxl

    rollups = _rollups(built)
    wb = openpyxl.load_workbook(built / "bom.xlsx")
    assert wb.sheetnames == ["BOM", "Rollups", "Parts", "Manufacturer"]

    ws = wb["BOM"]
    header = [c.value for c in ws[1]]
    assert header == COLUMNS
    assert all(c.font.bold for c in ws[1])
    assert ws.freeze_panes == "A2"
    values = {
        str(ws.cell(r, 1).value): ws.cell(r, 2).value
        for r in range(2, ws.max_row + 1)
        if ws.cell(r, 1).value is not None
    }
    assert values["total_cost_usd"] == rollups["total_cost"]
    assert values["total_mass_kg"] == rollups["total_mass"]
    assert values["part_count"] == rollups["part_count"]
    assert not any(
        isinstance(c.value, str) and c.value.startswith("=") for row in ws.iter_rows() for c in row
    )
    assert SOURCING_NOTE in [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]

    ids = [ws.cell(r, 3).value for r in range(2, ws.max_row + 1)]
    bracket_row = ids.index("bracket") + 2
    ext_col = COLUMNS.index("ext_cost_usd") + 1
    assert ws.cell(bracket_row, ext_col).value == pytest.approx(8.0)
    assert ws.cell(bracket_row, ext_col).number_format == "0.00"
    assert ws.cell(bracket_row, COLUMNS.index("mass_kg") + 1).number_format == "0.000"

    rs = {r[0]: r[1] for r in wb["Rollups"].iter_rows(min_row=2, values_only=True)}
    assert rs["total_cost"] == rollups["total_cost"]
    assert rs["part_count"] == rollups["part_count"]

    parts = [r[0] for r in wb["Parts"].iter_rows(min_row=2, values_only=True)]
    assert parts == list(rollups["parts"])

    base_text = next(
        ln["manufacturer"]
        for top in rollups["top_assemblies"]
        for ln in rollups["structure"][top]
        if ln["id"] == "base"
    )
    man = {r[0]: r[4] for r in wb["Manufacturer"].iter_rows(min_row=2, values_only=True)}
    assert man["part://base"] == base_text
    assert man["part://m4x10-screw"].startswith("buy ")
