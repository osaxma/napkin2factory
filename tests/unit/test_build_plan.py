"""v0.4 M3a: build units and phantoms (P7), the derived build plan (P8), hardware counted (P9).

The ``units`` fixture: the top (a unit) holds the base, a phantom cover group (cover, stop,
two nuts) and the pod, a unit of its own (plate, two rails).  Joints: pod-mount (base to
plate, a port of the pod: a standoff and two bolts per hole), rail-plate (inside the pod,
two holes per rail), cover-base (``fastener:`` alias, one bolt per hole) and stop-cover,
which shares cover-base's holes and lists only its nuts (S-105's shared-hole rule).
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from pydantic import ValidationError
from selectolax.lexbor import LexborHTMLParser

from napkin2factory.bom.build_plan import derive_build_plan, step_markdown
from napkin2factory.core import lint as lint_mod
from napkin2factory.core.build_units import BuildUnits
from napkin2factory.core.checks.hardware import hardware_checks
from napkin2factory.core.graph.runner import run_build
from napkin2factory.core.graph.values import unvalued_reads
from napkin2factory.core.loader import Doc, Product, load_product
from napkin2factory.core.model import DocMeta, Interface
from napkin2factory.docs.engine import DocRenderError, render_doc
from tests.conftest import EXAMPLES, FIXTURES, copy_product
from tests.validate_outputs_full import validate_full

UNITS = FIXTURES / "units"


def _copy(tmp_path: Path) -> Path:
    return copy_product(UNITS, tmp_path / "units", git=False)


def _edit(root: Path, rel: str, old: str, new: str) -> Product:
    path = root / "model" / rel
    text = path.read_text()
    assert old in text, (rel, old)
    path.write_text(text.replace(old, new))
    return load_product(root)


def _rows(product: Product) -> dict[str, tuple[str, str, dict]]:  # type: ignore[type-arg]
    return {r.id: (r.status, r.message, r.detail) for r in hardware_checks(product)}


# ---- the plan ------------------------------------------------------------------------


def test_plan_order_children_first_phantoms_folded() -> None:
    product = load_product(UNITS)
    plan = derive_build_plan(product)
    assert plan["order"] == ["pod", "top"]  # the phantom has no step
    assert "frame" not in plan["units"]
    pod, top = plan["units"]["pod"], plan["units"]["top"]
    assert (pod["step"], top["step"]) == (1, 2)
    assert pod["driver"] == ["built-as-unit", "tested-as-unit"] and pod["parents"] == ["asm://top"]
    assert top["folded"] == [
        {
            "asm": "asm://frame",
            "name": "Cover group",
            "qty": 1,
            "words": "not put together on its own",
            "items": [5, 6, 7],  # gate D: which lines are the grouping's
        }
    ]
    items = [(it["item"], it["uri"], it["qty"], it["via"]) for it in top["items"]]
    assert items == [
        (1, "part://base", 1, "asm://top"),
        (2, "asm://pod", 1, "asm://top"),
        (3, "part://bolt", 10, "asm://top"),
        (4, "part://standoff", 4, "asm://top"),
        (5, "part://cover", 1, "asm://frame"),  # the phantom's lines, folded in
        (6, "part://stop", 1, "asm://frame"),
        (7, "part://nut", 2, "asm://frame"),
    ]
    assert top["items"][1]["kind"] == "unit" and top["items"][1]["step"] == 1
    # joints close where both parties' homes meet; the phantom's own joint folds up
    closes = {j["interface"]: (j["closes_in"], j["between_units"]) for j in top["joints"]}
    assert closes == {
        "iface://cover-base": ("asm://top", False),
        "iface://pod-mount": ("asm://top", True),  # a joint between units (a port of the pod)
        "iface://stop-cover": ("asm://frame", False),
    }
    (rail,) = pod["joints"]
    assert rail["occurrences"] == 2 and rail["hardware"][0]["qty"] == 4  # two rails, 2 holes
    # static check ids; the report joins verdicts
    assert top["checks"][:3] == ["hardware://top", "count://top", "interference://top"]
    assert "count://frame" in top["checks"] and "iface://pod-mount" in top["checks"]


def test_a_phantom_inside_a_phantom_multiplies_through(tmp_path: Path) -> None:
    root = _copy(tmp_path)
    (root / "model/assemblies/outer.yaml").write_text(
        "id: outer\nrev: A\nname: Outer group\nbuild: phantom\nports: [iface://cover-base]\n"
        "children:\n  - {asm: asm://frame, qty: 1}\n"
    )
    product = _edit(
        root, "assemblies/top.yaml", "{asm: asm://frame, qty: 1}", "{asm: asm://outer, qty: 1}"
    )
    units = BuildUnits(product)
    assert units.order() == ["pod", "top"]
    assert units.folded("top") == [("top", 1), ("outer", 1), ("frame", 1)]
    assert [it.uri for it in units.items("top")][-3:] == [
        "part://cover",
        "part://stop",
        "part://nut",
    ]


def test_flat_examples_have_one_step() -> None:
    for name in ("bracket", "enclosure"):
        plan = derive_build_plan(load_product(EXAMPLES / name))
        assert plan["order"] == [name]
        assert plan["units"][name]["title"].endswith("the final assembly")


# ---- hardware (P9) -------------------------------------------------------------------


def test_hardware_counts_pass_on_the_fixture() -> None:
    rows = _rows(load_product(UNITS))
    assert {k: v[0] for k, v in rows.items()} == {
        "hardware://pod": "pass",
        "hardware://top": "pass",
    }
    top = rows["hardware://top"][2]["parts"]
    assert top["part://bolt"] == {
        "declared": 10,
        "demand": 10,
        "interfaces": ["cover-base", "pod-mount"],
    }
    assert rows["hardware://top"][2]["folded"] == ["asm://frame"]


def test_shared_holes_count_once() -> None:
    """stop-cover shares cover-base's holes and lists only its nuts: no second set of bolts."""
    units = BuildUnits(load_product(UNITS))
    demand = units.demand("top")
    assert demand["part://bolt"][0] == 8 + 2  # pod-mount 4 × 2, cover-base 2 × 1, once
    assert demand["part://nut"] == (2, ["stop-cover"])


def test_fastener_alias_reads_as_one_per_hole() -> None:
    iface = load_product(UNITS).interfaces["cover-base"]
    assert [(str(h.part), h.per_hole) for h in iface.hardware_items()] == [("part://bolt", 1)]
    data = iface.model_dump(mode="json")
    explicit = Interface.model_validate(
        {**data, "spec": {**data["spec"], "hardware": []}}
    )  # `hardware: []` says "nothing here", whatever `fastener:` says
    assert explicit.hardware_items() == []


def test_a_missing_bolt_fails_naming_joints_and_number(tmp_path: Path) -> None:
    product = _edit(
        _copy(tmp_path),
        "assemblies/top.yaml",
        "{part: part://bolt, qty: 10}",
        "{part: part://bolt, qty: 9}",
    )
    status, message, detail = _rows(product)["hardware://top"]
    assert status == "fail"
    assert "M6 x 16 bolt: 1 missing, 9 listed where its joints need 10" in message
    # gate D: the owner reads what to do, not only what is wrong
    assert (
        "Add 1 to its line in the parts of Top, or correct the joint's hardware if fewer are "
        "needed" in message
    )
    assert "base to cover" in message and "base to plate" in message
    assert detail["parts"]["part://bolt"]["missing"] == 1
    assert detail["parts"]["part://bolt"]["interfaces"] == ["cover-base", "pod-mount"]


def test_surplus_warns_without_spare_and_passes_with_one(tmp_path: Path) -> None:
    root = _copy(tmp_path)
    product = _edit(
        root, "assemblies/pod.yaml", "{part: part://bolt, qty: 4}", "{part: part://bolt, qty: 6}"
    )
    status, message, _ = _rows(product)["hardware://pod"]
    assert status == "warn" and "2 more than needed, 6 listed where its joints need 4" in message
    product = _edit(
        root,
        "assemblies/pod.yaml",
        "{part: part://bolt, qty: 6}",
        '{part: part://bolt, qty: 6, spare: "two for the installer"}',
    )
    status, message, detail = _rows(product)["hardware://pod"]
    assert status == "pass" and "2 spare (two for the installer)" in message
    assert detail["parts"]["part://bolt"]["spare"] == ["two for the installer"]


def test_hardware_in_the_wrong_unit_fails_there(tmp_path: Path) -> None:
    """The pod-mount bolts belong where the joint closes (the top), not in the pod."""
    root = _copy(tmp_path)
    _edit(
        root, "assemblies/top.yaml", "{part: part://bolt, qty: 10}", "{part: part://bolt, qty: 2}"
    )
    product = _edit(
        root, "assemblies/pod.yaml", "{part: part://bolt, qty: 4}", "{part: part://bolt, qty: 12}"
    )
    rows = _rows(product)
    assert rows["hardware://top"][0] == "fail" and rows["hardware://pod"][0] == "warn"


def test_hardware_schema() -> None:
    base = {"id": "x", "parties": ["part://a", "part://b"]}
    with pytest.raises(ValidationError, match="exactly one of per_hole"):
        Interface.model_validate(
            {
                **base,
                "kind": "bolt-pattern",
                "spec": {"count": 2, "diameter": "6 mm", "hardware": [{"part": "part://a"}]},
            }
        )
    with pytest.raises(ValidationError, match="a mating face has no holes"):
        Interface.model_validate(
            {
                **base,
                "kind": "mating-face",
                "spec": {
                    "width": "10 mm",
                    "depth": "10 mm",
                    "hardware": [{"part": "part://r", "per_hole": 1}],
                },
            }
        )
    face = Interface.model_validate(
        {
            **base,
            "kind": "mating-face",
            "spec": {
                "width": "10 mm",
                "depth": "10 mm",
                "hardware": [{"part": "part://r", "per_joint": 4}],
            },
        }
    )
    assert face.holes == 0 and face.hardware_items()[0].per_joint == 4


@pytest.mark.parametrize("name", ["bracket", "enclosure"])
def test_v03_files_dump_as_before(name: str) -> None:
    """P14: no new key appears in a dump (sign-off evidence keys hash these dumps).

    The two flat examples have the v0.3 shape; the kiln uses the new fields since M3b.
    """
    product = load_product(EXAMPLES / name)
    for iface in product.interfaces.values():
        assert "hardware" not in iface.model_dump(mode="json")["spec"]
    for asm in product.assemblies.values():
        data = asm.model_dump(mode="json")
        assert "build" not in data and "driver" not in data
        assert all(set(c) <= {"part", "asm", "qty"} for c in data["children"])


def test_kiln_hardware_is_all_accounted_for() -> None:
    """M3b (S-105 option A): every unit's joint hardware matches its lines exactly."""
    product = load_product(EXAMPLES / "kiln")
    rows = _rows(product)
    assert {k: v[0] for k, v in rows.items()} == {
        "hardware://controls": "pass",
        "hardware://door": "pass",
        "hardware://kiln": "pass",
    }
    kiln = rows["hardware://kiln"][2]["parts"]
    want = {
        "part://m6x16-bolt": 38,  # 24 standoff walls + 4 post + 2 keeper + 8 controls
        "part://m6x50-bolt": 4,  # one through bolt per bottom corner
        "part://m6-standoff-25": 16,  # the bottom four listed once (shared holes)
        "part://m6-standoff-40": 4,
        "part://m6-nut": 6,
        "part://m6-rivnut": 4,
    }
    assert {k: (v["declared"], v["demand"]) for k, v in kiln.items()} == {
        k: (n, n) for k, n in want.items()
    }
    assert rows["hardware://controls"][2]["parts"]["part://lid-screw"]["demand"] == 4
    # the lining is folded into the final assembly; two hinges on one post share their count
    units = BuildUnits(product)
    assert units.order() == ["controls", "door", "kiln"]
    assert units.folded_ids("kiln") == ["kiln", "lining"]
    joints = {j.iface: j for j in units.joints("kiln")}
    assert joints["hinge-lower"].occurrences == joints["hinge-upper"].occurrences == 1
    assert joints["hinge-lower"].between_units and joints["controller-mount"].between_units
    assert not joints["guard-hinge-post"].between_units


# ---- lint (P7) -------------------------------------------------------------------------


def test_missing_driver_lint_warns(tmp_path: Path) -> None:
    root = _copy(tmp_path)
    assert not [p for p in lint_mod.lint(load_product(root)).items if p.rule.startswith("L13")]
    product = _edit(root, "assemblies/pod.yaml", "driver: [built-as-unit, tested-as-unit]\n", "")
    (finding,) = [p for p in lint_mod.lint(product).items if p.rule == "L13-assembly-driver"]
    assert finding.severity == "warn" and finding.subject == "asm://pod"
    assert "nothing says why" in finding.message and "build: phantom" in finding.hint
    product = _edit(root, "assemblies/top.yaml", "name: Top\n", "name: Top\nbuild: phantom\n")
    subjects = {p.subject for p in lint_mod.lint(product).items if p.rule == "L13-assembly-driver"}
    assert subjects == {"asm://pod", "asm://top"}


def test_unknown_driver_is_refused() -> None:
    from napkin2factory.core.model import Assembly

    with pytest.raises(ValidationError):
        Assembly.model_validate({"id": "a", "name": "A", "driver": ["because"]})


# ---- docs helper (D5) ------------------------------------------------------------------


def test_build_step_records_its_read(tmp_path: Path) -> None:
    product = load_product(UNITS)
    doc = Doc(
        meta=DocMeta(id="assembly", title="Assembly"),
        body='{{ build_step("asm://pod") }}',
        path=tmp_path / "assembly.md",
    )
    rendered = render_doc(product, doc)
    assert "bom://build-plan#units.pod" in rendered.reads
    plan = derive_build_plan(product)
    assert rendered.markdown == step_markdown(plan, "pod")
    assert "| 3 | M6 x 16 bolt | 4 | bought |" in rendered.markdown
    assert "2 × plate to rail, 2 holes each: 4 × M6 x 16 bolt" in rendered.markdown
    phantom = Doc(meta=doc.meta, body='{{ build_step("asm://frame") }}', path=doc.path)
    with pytest.raises(DocRenderError, match="grouping"):
        render_doc(product, phantom)


# ---- built outputs ---------------------------------------------------------------------


@pytest.fixture(scope="module")
def built(tmp_path_factory: pytest.TempPathFactory) -> Path:
    root = copy_product(UNITS, tmp_path_factory.mktemp("units") / "units", git=False)
    run_build(load_product(root), quiet=True, render=False)
    return root


def test_plan_file_xlsx_sheet_and_one_assembly_sheet_per_unit(built: Path) -> None:
    assert (built / "build/bom/build-plan.json").is_file()
    wb = openpyxl.load_workbook(built / "build/bom/bom.xlsx", data_only=True)
    assert "Build plan" in wb.sheetnames
    rows = [tuple(r) for r in wb["Build plan"].iter_rows(values_only=True)]
    assert rows[0] == ("step", "unit", "row", "item", "uri", "name", "qty", "note")
    kinds = [(r[0], r[2]) for r in rows[1:]]
    assert kinds[0] == (1, "step") and (2, "step") in kinds and (2, "joint") in kinds
    assert (2, 3, "part://bolt", 10) in {(r[0], r[3], r[4], r[6]) for r in rows[1:]}
    sheets = sorted(p.name for p in (built / "build/drawings/assemblies").glob("*.svg"))
    assert sheets == ["pod.svg", "top.svg"]  # none for the phantom
    svg = (built / "build/drawings/assemblies/top.svg").read_text()
    assert "Step 2: Top, the final assembly" in svg and "4 × item 4 + 8 × item 3" in svg
    assert "no picture" in svg  # built without renders
    problems: list[str] = []
    validate_full(built, load_product(built), problems)
    assert not [p for p in problems if "build" in p or "sheet" in p or "step" in p], problems


def test_report_build_plan_section(built: Path) -> None:
    html = (built / "build/report/index.html").read_text(encoding="utf-8")
    tree = LexborHTMLParser(html)
    section = tree.css_first("#build")
    assert section is not None
    text = section.text()
    assert "Step 1: Pod, built as one unit" in text and "Step 2: Top, the final assembly" in text
    # gate D: the grouping explained, its lines marked; the reason said with "and"
    assert (
        "The cover group (items 5–7) is not put together on its own: lay it in place during "
        "this step." in text
    )
    assert "a joint between units" in text
    assert tree.css_first("#step-pod svg") is not None  # the sheet, inlined
    assert (
        "Built on its own because it is put together and fitted as one piece and tested as "
        "one unit." in text
    )
    assert "Fitted in step 2 (Top, the final assembly)." in " ".join(text.split())
    notes = [" ".join(r.text().split()) for r in tree.css("#step-top tbody tr")]
    assert sum("cover group" in n for n in notes) == 3
    # P13 (M6): the joints between units, each with its units, parts, hardware and verdict
    joints = tree.css_first("#joints")
    assert joints is not None and tree.css_first('nav a[data-sec="joints"]') is not None
    (row,) = joints.css("tbody tr")
    words = " ".join(row.text().split())
    assert row.attributes["id"] == "joint-pod-mount"
    assert "Top and Pod" in words and "fitted in step 2" in words and "bolted, 4 holes" in words
    assert "4 × M6 standoff 25" in words and "8 × M6 x 16 bolt" in words
    assert row.css_first(".badge") is not None


def test_build_plan_is_rebuilt_when_hardware_changes(built: Path) -> None:
    from napkin2factory.core.graph.runner import explain_build

    _edit(built, "interfaces/rail-plate.yaml", "per_hole: 1", "per_hole: 2")
    plan = {rid: run for rid, run, _ in explain_build(load_product(built), render=False)}
    assert plan["bom"] and plan["asmsheet:pod"] and plan["checks"]


# ---- review gate C: the plan is a recorded read with a value (B1, B2, C7; S-117) ------------

_GEOMETRIC = ("part:", "asm:", "drawing:")


@pytest.fixture
def built_with_page(tmp_path: Path) -> Path:
    """The units fixture with an assembly page that prints the pod's step, built once."""
    root = _copy(tmp_path)
    page = root / "model/docs/assembly.md"
    page.parent.mkdir(parents=True, exist_ok=True)
    page.write_text(
        '---\nid: assembly\ntitle: Assembly\n---\n\n# Assembly\n\n{{ build_step("asm://pod") }}\n'
    )
    run_build(load_product(root), quiet=True, render=False)
    return root


def test_every_recorded_read_has_a_value(built_with_page: Path) -> None:
    """C7: no key in the manifest is valued ``unknown:`` (the class behind B1)."""
    assert unvalued_reads(load_product(built_with_page), built_with_page / "build") == {}


def test_a_plan_only_change_rebuilds_the_sheet_the_page_and_the_report(
    built_with_page: Path,
) -> None:
    """B1: a driver change is in the plan only: the pod's sheet, the page and the report follow."""
    from napkin2factory.core.checks.engine import doc_staleness
    from napkin2factory.core.graph.runner import explain_build

    root = built_with_page
    product = _edit(
        root,
        "assemblies/pod.yaml",
        "[built-as-unit, tested-as-unit]",
        "[built-as-unit, tested-as-unit, varies]",
    )
    predicted = {rid for rid, run, _ in explain_build(product, render=False) if run}
    assert {"bom", "asmsheet:pod", "doc:assembly", "report"} <= predicted, predicted
    run_build(product, quiet=True, render=False, target="bom")
    stale = [r for r in doc_staleness(product, root / "build") if r.status == "fail"]
    assert [r.id for r in stale] == ["doc://assembly"]
    assert "bom://build-plan#units.pod" in stale[0].message
    ran = run_build(product, quiet=True, render=False)["ran"]
    assert {"asmsheet:pod", "doc:assembly", "report"} <= set(ran), ran
    assert not [r for r in ran if r.startswith(_GEOMETRIC)], ran
    assert "asmsheet:top" not in ran  # the top's step did not change
    assert "comes in variants" in (root / "build/drawings/assemblies/pod.svg").read_text()
    assert "comes in variants" in (root / "build/docs/assembly/page.md").read_text()
    assert "comes in variants" in (root / "build/report/index.html").read_text()
    assert [r.status for r in doc_staleness(load_product(root), root / "build")] == ["pass"]


def test_a_renamed_part_reaches_its_units_sheet(built_with_page: Path) -> None:
    """B1: the gate renamed a door part; the sheet kept the old name."""
    root = built_with_page
    product = _edit(root, "parts/plate.yaml", "name: plate", "name: Plate RENAMED")
    ran = run_build(product, quiet=True, render=False)["ran"]
    assert {"asmsheet:pod", "doc:assembly", "report"} <= set(ran), ran
    assert "Plate RENAMED" in (root / "build/drawings/assemblies/pod.svg").read_text()
    assert "Plate RENAMED" in (root / "build/docs/assembly/page.md").read_text()


def test_the_hardware_verdict_follows_spare_without_a_forced_rerun(tmp_path: Path) -> None:
    """B2: add then remove ``spare:``; the hardware row warns, passes, warns again."""
    import json

    root = _copy(tmp_path)
    _edit(
        root, "assemblies/top.yaml", "{part: part://bolt, qty: 10}", "{part: part://bolt, qty: 12}"
    )

    def row() -> tuple[str, str]:
        run_build(load_product(root), quiet=True, render=False, target="checks")
        data = json.loads((root / "build/checks/checks.json").read_text())
        r = next(r for r in data["results"] if r["id"] == "hardware://top")
        return r["status"], r["message"]

    assert row()[0] == "warn"
    line = "{part: part://bolt, qty: 12}"
    _edit(root, "assemblies/top.yaml", line, line[:-1] + ', spare: "two for the cover"}')
    status, message = row()
    assert status == "pass" and "two for the cover" in message
    _edit(root, "assemblies/top.yaml", ', spare: "two for the cover"}', "}")
    status, message = row()
    assert status == "warn" and "two for the cover" not in message


# ---- gate D (S-118): joint names, build order, a joint closed later, the sheet's text --


def test_kiln_joints_are_named_and_in_build_order() -> None:
    plan = derive_build_plan(load_product(EXAMPLES / "kiln"))
    words = [j["words"] for j in plan["units"]["kiln"]["joints"]]
    assert words == [
        "Shell to guard, bottom wall (the stand's bolts pass through)",
        "Shell to guard, left wall",
        "Shell to guard, right wall",
        "Shell to guard, top wall",
        "Body onto the stand",
        "Hinge post onto the guard",
        "Latch keeper onto the guard",
        "Door to kiln, lower hinge (welded)",
        "Door to kiln, upper hinge (welded)",
        "Door to kiln, draw latch",
        "Controls onto the guard's right wall",
        "Lid onto the controller box",  # closed later, carried from step 1
    ]
    assert len(set(words)) == len(words)  # every joint can be told apart


def test_unnamed_joint_reads_as_its_two_parts() -> None:
    plan = derive_build_plan(load_product(UNITS))
    assert {j["words"] for j in plan["units"]["top"]["joints"]} >= {"base to cover"}


def test_lid_joint_is_closed_later_in_the_final_assembly() -> None:
    """Gate D, concern 13: the controls' lid goes on after the wiring, as the manual says."""
    product = load_product(EXAMPLES / "kiln")
    assert BuildUnits(product).closed_later("controller-lid-seat") == "kiln"
    plan = derive_build_plan(product)
    (lid,) = plan["units"]["controls"]["joints"]
    assert lid["closed_later_in"] == "asm://kiln" and lid["closed_later_step"] == 3
    carried = plan["units"]["kiln"]["joints"][-1]
    assert carried["interface"] == "iface://controller-lid-seat" and carried["from_step"] == 1
    assert "closed_later_in" not in carried
    assert carried not in plan["between_units"]
    controls = step_markdown(plan, "controls")
    assert "- Lid onto the controller box (closed later, in step 3): 4 × Self-tapping" in controls
    kiln = step_markdown(plan, "kiln")
    assert "Lid onto the controller box (left open in step 1): 4 × Self-tapping" in kiln
    assert "(listed in step 1)" in kiln
    # its hardware is counted once, where the screws are listed
    rows = _rows(product)
    assert rows["hardware://controls"][0] == rows["hardware://kiln"][0] == "pass"
    assert "part://lid-screw" not in rows["hardware://kiln"][2]["parts"]


def test_closed_in_naming_no_later_unit_is_ignored_and_linted(tmp_path: Path) -> None:
    root = _copy(tmp_path)
    product = _edit(
        root,
        "interfaces/rail-plate.yaml",
        "kind: bolt-pattern\n",
        "kind: bolt-pattern\nclosed_in: asm://frame\n",
    )
    assert BuildUnits(product).closed_later("rail-plate") is None  # a phantom: not a step
    (finding,) = [p for p in lint_mod.lint(product).items if p.rule == "L14-closed-in"]
    assert finding.severity == "warn" and finding.subject == "iface://rail-plate"
    product = _edit(root, "interfaces/rail-plate.yaml", "asm://frame", "asm://top")
    assert BuildUnits(product).closed_later("rail-plate") == "top"
    assert not [p for p in lint_mod.lint(product).items if p.rule == "L14-closed-in"]


def test_interface_name_is_a_label_not_evidence(tmp_path: Path) -> None:
    from napkin2factory.core.checks.evidence import evidence_key

    root = _copy(tmp_path)
    before = evidence_key(load_product(root), root / "build", "iface://rail-plate")
    product = _edit(
        root,
        "interfaces/rail-plate.yaml",
        "kind: bolt-pattern\n",
        "kind: bolt-pattern\nname: Rails onto the plate\n",
    )
    assert evidence_key(product, root / "build", "iface://rail-plate") == before
    # unset, the new fields do not dump: v0.3 interfaces hash as before
    dumped = product.interfaces["cover-base"].model_dump()
    assert "name" not in dumped and "closed_in" not in dumped


def test_sheet_wraps_long_cells_instead_of_shrinking_below_2_2_mm() -> None:
    import re

    from napkin2factory.drawings.assembly_sheet import MIN_TEXT, compose_sheet

    plan = derive_build_plan(load_product(EXAMPLES / "kiln"))
    step = plan["units"]["kiln"]
    long = " ".join(["A very long joint name"] * 4)  # far wider than its column
    step["joints"][0]["words"] = long
    svg, _stable, _meta = compose_sheet(step, 3, [], "unsnapshotted")
    tables = svg.split('<g id="tables">', 1)[1].split("</g>", 1)[0]
    sizes = [float(x) for x in re.findall(r'font-size="([0-9.]+)"', tables)]
    assert min(sizes) >= MIN_TEXT
    assert f">{long}<" not in tables and ">A very long joint name A very" in tables
    assert "not yet saved" in svg and "unsnapshotted" not in svg
    assert "PARTS FOR THIS STEP" in svg and "BOM LINE" not in svg and "EXPLODE" not in svg
    assert "not numbered" in svg and "and the BOM" not in svg


def test_each_sheet_records_the_plans_order(built: Path) -> None:
    """Gate C's B-0068: "STEP n of N" reads the order, so a unit added redraws every sheet."""
    from napkin2factory.core.graph.manifest import Manifest

    entries = Manifest.load(built / "build").entries
    for uid in ("pod", "top"):
        entry = entries[f"asmsheet:{uid}"]
        assert "bom://build-plan#order" in entry.reads
        assert not entry.values["bom://build-plan#order"].startswith("unknown:")
