"""BOM files: validation (D15), flat CSV and a formatted XLSX workbook (D18, D17).

The rollup numbers are written as values, never formulas, so readers that do not
calculate (CSV importers, ``openpyxl`` with ``data_only``) see the same totals as
``rollups.json``.
"""

from __future__ import annotations

import csv
import json
from collections.abc import Iterator
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from napkin2factory.core.graph.rules import RunContext

COLUMNS = [
    "level",
    "uri",
    "id",
    "rev",
    "name",
    "kind",
    "qty",
    "total_qty",
    "process",
    "material",
    "finish",
    "tolerance_class",
    "unit_cost_usd",
    "ext_cost_usd",
    "lead_time_days",
    "mass_kg",
    "sourcing_status",
    "source",
    "manufacturer",
]

TOTALS = [
    ("total_cost_usd", "total_cost"),
    ("total_mass_kg", "total_mass"),
    ("part_count", "part_count"),
    ("custom_count", "custom_count"),
    ("custom_ratio", "custom_ratio"),
    ("max_lead_time_days", "max_lead_time"),
    ("cost_complete", "cost_complete"),
    ("mass_complete", "mass_complete"),
]

SOURCING_NOTE = (
    "Unverified sources are marked in sourcing_status; costs are declared estimates unless "
    "basis says quote"
)

_WIDTHS = {
    "level": 6,
    "uri": 24,
    "id": 16,
    "rev": 6,
    "name": 32,
    "kind": 10,
    "qty": 6,
    "total_qty": 9,
    "process": 16,
    "material": 22,
    "finish": 18,
    "tolerance_class": 14,
    "unit_cost_usd": 13,
    "ext_cost_usd": 13,
    "lead_time_days": 14,
    "mass_kg": 10,
    "sourcing_status": 15,
    "source": 28,
    "manufacturer": 60,
}


def iter_lines(rollups: dict[str, Any]) -> Iterator[dict[str, Any]]:
    """Every BOM line, depth-first through sub-assemblies, for each top assembly.

    Args:
        rollups: The ``rollups.json`` payload.

    Yields:
        Line dicts (as in ``rollups["structure"]``) in display order.
    """

    def walk(lines: list[dict[str, Any]]) -> Iterator[dict[str, Any]]:
        for ln in lines:
            yield ln
            yield from walk(ln.get("children") or [])

    for top in rollups.get("top_assemblies", []):
        yield from walk(rollups.get("structure", {}).get(top, []))


def line_row(ln: dict[str, Any]) -> dict[str, Any]:
    """Map a rollup line dict onto the :data:`COLUMNS` (``None`` for unknown numbers).

    Args:
        ln: One line dict from ``rollups["structure"]``.

    Returns:
        A dict keyed by the BOM column names.
    """
    return {
        "level": ln.get("level", 0),
        "uri": ln.get("uri", ""),
        "id": ln.get("id", ""),
        "rev": ln.get("rev", ""),
        "name": ln.get("name", ""),
        "kind": ln.get("kind", ""),
        "qty": ln.get("qty"),
        "total_qty": ln.get("total_qty"),
        "process": ln.get("process", ""),
        "material": ln.get("material", ""),
        "finish": ln.get("finish", ""),
        "tolerance_class": ln.get("tolerance_class", ""),
        "unit_cost_usd": ln.get("unit_cost"),
        "ext_cost_usd": None if ln.get("ext_cost") is None else round(ln["ext_cost"], 6),
        "lead_time_days": ln.get("lead_time_days"),
        "mass_kg": ln.get("mass_kg"),
        "sourcing_status": ln.get("sourcing_status", ""),
        "source": ln.get("source", ""),
        "manufacturer": ln.get("manufacturer", ""),
    }


def _csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


def write_csv(path: Path, rollups: dict[str, Any]) -> Path:
    """Write the flat ``bom.csv`` (one row per line, then a blank row and the totals).

    Args:
        path: Destination file.
        rollups: The ``rollups.json`` payload.

    Returns:
        ``path``.
    """
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(COLUMNS)
        for ln in iter_lines(rollups):
            row = line_row(ln)
            writer.writerow([_csv_value(row[c]) for c in COLUMNS])
        writer.writerow([])
        for label, key in TOTALS:
            writer.writerow([label, _csv_value(rollups.get(key))])
        writer.writerow(
            [
                "note",
                "This design was produced with an AI system and has not been reviewed "
                "by a licensed engineer.",
            ]
        )
    return path


def write_xlsx(path: Path, rollups: dict[str, Any]) -> Path:
    """Write ``bom.xlsx`` with the BOM, Rollups, Parts and Manufacturer sheets.

    Args:
        path: Destination file.
        rollups: The ``rollups.json`` payload.

    Returns:
        ``path``.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = bold
    money = {COLUMNS.index(c) + 1 for c in ("unit_cost_usd", "ext_cost_usd")}
    mass_col = COLUMNS.index("mass_kg") + 1
    name_col = COLUMNS.index("name") + 1
    for ln in iter_lines(rollups):
        row = line_row(ln)
        row["name"] = "  " * int(row["level"] or 0) + str(row["name"])
        ws.append([row[c] for c in COLUMNS])
        r = ws.max_row
        for col in money:
            ws.cell(r, col).number_format = "0.00"
        ws.cell(r, mass_col).number_format = "0.000"
        ws.cell(r, name_col).alignment = Alignment(indent=int(row["level"] or 0))
    ws.freeze_panes = "A2"
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _WIDTHS.get(col, 12)

    ws.append([])
    ws.append(["Totals"])
    ws.cell(ws.max_row, 1).font = bold
    for label, key in TOTALS:
        ws.append([label, rollups.get(key)])
        r = ws.max_row
        ws.cell(r, 1).font = bold
        if key == "total_cost":
            ws.cell(r, 2).number_format = "0.00"
        elif key == "total_mass":
            ws.cell(r, 2).number_format = "0.000"
    ws.append([])
    ws.append([SOURCING_NOTE])
    ws.cell(ws.max_row, 1).font = Font(italic=True)

    rs = wb.create_sheet("Rollups")
    rs.append(["key", "value"])
    for cell in rs[1]:
        cell.font = bold
    for key in (
        "total_cost",
        "total_cost_unit",
        "cost_complete",
        "total_mass",
        "total_mass_unit",
        "mass_complete",
        "part_count",
        "distinct_parts",
        "custom_count",
        "custom_ratio",
        "max_lead_time",
        "max_lead_time_unit",
    ):
        if key in rollups:
            rs.append([key, rollups[key]])
    rs.append(["top_assemblies", ", ".join(rollups.get("top_assemblies", []))])
    rs.append(["orphans", ", ".join(rollups.get("orphans", []))])
    rs.column_dimensions["A"].width = 20
    rs.column_dimensions["B"].width = 30

    ps = wb.create_sheet("Parts")
    part_cols = [
        "uri",
        "name",
        "kind",
        "total_qty",
        "unit_cost",
        "lead_time_days",
        "mass_kg",
        "sourcing_status",
    ]
    ps.append(part_cols)
    for cell in ps[1]:
        cell.font = bold
    for part in rollups.get("parts", {}).values():
        ps.append([part.get(c) for c in part_cols])
        r = ps.max_row
        ps.cell(r, part_cols.index("unit_cost") + 1).number_format = "0.00"
        ps.cell(r, part_cols.index("mass_kg") + 1).number_format = "0.000"
    for i, w in enumerate((24, 32, 10, 9, 10, 14, 10, 15), start=1):
        ps.column_dimensions[get_column_letter(i)].width = w
    ps.freeze_panes = "A2"

    ms = wb.create_sheet("Manufacturer")
    ms.append(["uri", "name", "kind", "total_qty", "instruction"])
    for cell in ms[1]:
        cell.font = bold
    seen: set[str] = set()
    for ln in iter_lines(rollups):
        if ln.get("kind") == "assembly" or ln.get("uri") in seen:
            continue
        seen.add(str(ln.get("uri")))
        total = (
            rollups.get("parts", {}).get(ln.get("uri"), {}).get("total_qty", ln.get("total_qty"))
        )
        ms.append(
            [ln.get("uri"), ln.get("name"), ln.get("kind"), total, ln.get("manufacturer", "")]
        )
        ms.cell(ms.max_row, 5).alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate((24, 32, 10, 9, 80), start=1):
        ms.column_dimensions[get_column_letter(i)].width = w
    ms.freeze_panes = "A2"

    wb.save(path)
    return path


def write_bom_files(ctx: RunContext, rollups: dict[str, Any]) -> list[Path]:
    """Write ``validation.json``, ``bom.csv`` and ``bom.xlsx``; return the files written.

    Args:
        ctx: The build context.
        rollups: The ``rollups.json`` payload from ``compute_rollups``.

    Returns:
        ``[validation.json, bom.csv, bom.xlsx]``.
    """
    from napkin2factory.bom.validation import validate_bom

    out_dir = ctx.build_dir / "bom"
    out_dir.mkdir(parents=True, exist_ok=True)
    validation = validate_bom(ctx.product, ctx.build_dir)
    path = out_dir / "validation.json"
    path.write_text(json.dumps(validation, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return [
        path,
        write_csv(out_dir / "bom.csv", rollups),
        write_xlsx(out_dir / "bom.xlsx", rollups),
    ]
